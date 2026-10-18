from unittest.mock import patch, MagicMock, AsyncMock
import pytest
from starlette.testclient import TestClient
from api.main import app
from api.deps import get_db
from services.dataset_ingestion import parse_uploaded_dataset, load_seed_datasets
from services.reference_retriever import retrieve_references, _score_item

client = TestClient(app)
AUTH_HEADERS = {"X-API-Key": "test-key-001"}


@pytest.fixture(autouse=True)
def override_db_dependency():
    """Mock the get_db dependency for fast unit tests without live Postgres."""
    mock_session = AsyncMock()
    mock_session.commit = AsyncMock()
    mock_session.add = MagicMock()
    mock_session.delete = AsyncMock()
    mock_session.execute = AsyncMock()

    mock_result = MagicMock()
    mock_result.all.return_value = []
    mock_result.scalars.return_value.all.return_value = []
    mock_result.scalar_one_or_none.return_value = None
    mock_session.execute.return_value = mock_result

    async def _mock_get_db():
        yield mock_session

    app.dependency_overrides[get_db] = _mock_get_db
    yield
    app.dependency_overrides.pop(get_db, None)


@pytest.mark.asyncio
async def test_seed_datasets_loading():
    """Verify that all seed JSON datasets load cleanly."""
    datasets = await load_seed_datasets()
    assert len(datasets) >= 4
    domains = {str(d.get("domain", "")) for d in datasets}
    assert any("FinTech" in d for d in domains)
    assert any("E-Commerce" in d for d in domains)
    assert any("Embedded" in d for d in domains)
    assert any("Distributed" in d or "Cloud" in d for d in domains)


def test_parse_uploaded_dataset_json():
    json_bytes = b'''
    [
        {
            "title": "Kafka Partition Lag Balancer",
            "topic": "Distributed Systems",
            "difficulty": "HARD",
            "scenario": "Balance consumer group lag across dynamic partitions.",
            "tags": ["kafka", "concurrency"]
        }
    ]
    '''
    items = parse_uploaded_dataset("kafka_cases.json", json_bytes)
    assert len(items) == 1
    assert items[0]["title"] == "Kafka Partition Lag Balancer"
    assert items[0]["topic"] == "Distributed Systems"
    assert items[0]["doc_type"] == "reference_question"


def test_parse_uploaded_dataset_csv():
    csv_bytes = b'''title,topic,difficulty,scenario,tags
Order Throttle,Rate Limiting,MEDIUM,Token bucket for order checkout,"rate_limit,ecommerce"
'''
    items = parse_uploaded_dataset("rates.csv", csv_bytes)
    assert len(items) == 1
    assert items[0]["title"] == "Order Throttle"
    assert "rate_limit" in items[0]["tags"]


def test_calculate_relevance():
    ref = {
        "title": "Idempotent Payment Processor",
        "topic": "Hash Maps",
        "subtopic": "Deduplication",
        "difficulty": "MEDIUM",
        "tags": ["fintech", "payments", "hash-table"],
    }
    score_match = _score_item(
        ref,
        topic="Hash Maps",
        subtopic="Deduplication",
        difficulty="MEDIUM",
    )
    score_unrelated = _score_item(
        ref,
        topic="Binary Trees",
        subtopic="Traversal",
        difficulty="EASY",
    )
    assert score_match > score_unrelated
    assert score_match >= 15


@pytest.mark.asyncio
async def test_retrieve_references_fallback():
    """When DB yields nothing or is unavailable, retriever falls back to pre-seeded files."""
    results = await retrieve_references(
        requirements={"topic": "Hash Maps", "difficulty": "MEDIUM", "subtopic": "Deduplication"},
        client_id="00000000-0000-0000-0000-000000000001",
        limit=3,
        db=None,
    )
    assert len(results) > 0
    assert any("Payment" in r.get("title", "") or "Ledger" in r.get("title", "") for r in results)


def test_api_list_datasets():
    res = client.get("/api/v1/datasets", headers=AUTH_HEADERS)
    assert res.status_code == 200
    data = res.json()
    assert isinstance(data, list)


def test_api_seed_datasets():
    res = client.post("/api/v1/datasets/seed", headers=AUTH_HEADERS)
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "SEEDED"
    assert data["count"] >= 4


def test_api_json_batch_upload():
    payload = {
        "domain": "Healthcare",
        "items": [
            {
                "title": "HL7 V2 Message Parser",
                "topic": "Parsing & Strings",
                "difficulty": "MEDIUM",
                "scenario": "Validate and parse segmented HL7 streams.",
                "tags": ["healthtech", "hl7"],
            }
        ],
    }
    res = client.post("/api/v1/datasets/json", json=payload, headers=AUTH_HEADERS)
    assert res.status_code == 201
    data = res.json()
    assert data["status"] == "INGESTED"
    assert data["count"] == 1


def test_parse_uploaded_dataset_docx():
    import io
    import docx

    doc = docx.Document()
    doc.add_heading("Fraud Velocity Rate Limiter", level=1)
    doc.add_paragraph("Topic: Sliding Window & Redis")
    doc.add_paragraph("Difficulty: HARD")
    doc.add_paragraph("Domain: FinTech & Banking")
    doc.add_paragraph("Tags: fraud, velocity, rate-limiting")
    doc.add_paragraph("Scenario: Build a sliding-window counter tracking suspicious card attempts across rolling 10-minute intervals.")

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    items = parse_uploaded_dataset("fraud_velocity.docx", docx_bytes)
    assert len(items) == 1
    item = items[0]
    assert "Fraud Velocity Rate Limiter" in item["title"]
    assert "Sliding Window" in item["topic"]
    assert item["difficulty"] == "HARD"
    assert "FinTech" in item["domain"]
    assert "fraud" in item["tags"]
    assert "sliding-window counter" in item["scenario"]


def test_parse_uploaded_dataset_pdf():
    pdf_data = b'''%PDF-1.4
1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj
2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj
3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >> endobj
4 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj
5 0 obj << /Length 95 >> stream
BT
/F1 12 Tf
72 712 Td
(Title: CAN-Bus Frame Ring Buffer\\nTopic: Circular Queue\\nDifficulty: HARD\\nScenario: Implement ring buffer) Tj
ET
endstream
endobj
xref
0 6
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000244 00000 n 
0000000318 00000 n 
trailer << /Size 6 /Root 1 0 R >>
startxref
465
%%EOF'''

    items = parse_uploaded_dataset("can_buffer.pdf", pdf_data)
    assert len(items) == 1
    assert "CAN-Bus Frame Ring Buffer" in items[0]["title"]
    assert items[0]["difficulty"] == "HARD"


def test_api_upload_docx_file():
    import io
    import docx

    doc = docx.Document()
    doc.add_heading("Warehouse AMR Robot Routing", level=1)
    doc.add_paragraph("Topic: Dijkstra Shortest Path")
    doc.add_paragraph("Difficulty: MEDIUM")
    doc.add_paragraph("Domain: E-Commerce & Logistics")
    doc.add_paragraph("Scenario: Route autonomous mobile robots in a dynamic grid with moving obstacles.")

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    files = {"file": ("amr_routing.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    data = {"doc_type": "DOMAIN_SCENARIO", "tags": "robotics,warehouse"}

    res = client.post("/api/v1/datasets/upload", files=files, data=data, headers=AUTH_HEADERS)
    assert res.status_code == 201
    res_data = res.json()
    assert res_data["status"] == "INGESTED"
    assert res_data["count"] == 1
    assert "AMR Robot Routing" in res_data["records"][0]["title"]
    assert res_data["records"][0]["format"] == "docx"


def test_api_upload_pdf_file():
    pdf_data = b'''%PDF-1.4
1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj
2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj
3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >> endobj
4 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj
5 0 obj << /Length 85 >> stream
BT
/F1 12 Tf
72 712 Td
(Title: Sensor Noise Filter\\nTopic: Signal Processing\\nDifficulty: EASY\\nScenario: Exponential filter) Tj
ET
endstream
endobj
xref
0 6
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000244 00000 n 
0000000318 00000 n 
trailer << /Size 6 /Root 1 0 R >>
startxref
455
%%EOF'''

    files = {"file": ("sensor_filter.pdf", pdf_data, "application/pdf")}
    data = {"doc_type": "DOMAIN_SCENARIO", "tags": "dsp,sensor"}

    res = client.post("/api/v1/datasets/upload", files=files, data=data, headers=AUTH_HEADERS)
    assert res.status_code == 201
    res_data = res.json()
    assert res_data["status"] == "INGESTED"
    assert res_data["count"] == 1
    assert "Sensor Noise Filter" in res_data["records"][0]["title"]
    assert res_data["records"][0]["format"] == "pdf"

